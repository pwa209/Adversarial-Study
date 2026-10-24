"""Synchronize the preserved proposal Word styles with the revised Markdown."""
from __future__ import annotations
import re
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def inline(paragraph, text):
    for token in re.split(r"(\*\*.*?\*\*|\[[^]]+\]\([^)]+\))", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("["):
            match = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", token)
            if match is None:
                paragraph.add_run(token)
                continue
            label, target = match.groups()
            if not target.startswith("https://"):
                paragraph.add_run(label)
                continue
            relationship = paragraph.part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), relationship)
            run = OxmlElement("w:r"); properties = OxmlElement("w:rPr")
            style = OxmlElement("w:rStyle"); style.set(qn("w:val"), "Hyperlink")
            properties.append(style); run.append(properties)
            literal = OxmlElement("w:t"); literal.text = label
            run.append(literal); link.append(run); paragraph._p.append(link)
        else:
            paragraph.add_run(token)


def update(root: Path):
    root = Path(root).resolve()
    path = root / "Adversarial_Consciousness_Research_Proposal.docx"
    history = root / "document_history"
    history.mkdir(exist_ok=True)
    backup = history / "Proposal_before_full_revision_20261004.docx"
    if not backup.exists():
        shutil.copy2(path, backup)
    document = Document(path)
    body = document._element.body
    for item in list(body):
        if item.tag != qn("w:sectPr"):
            body.remove(item)
    lines = (root / "Adversarial_Consciousness_Research_Proposal.md").read_text(encoding="utf-8").splitlines()
    index, paragraph_count, table_count = 0, 0, 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in cells):
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Normal Table"
            table.autofit = False
            widths = [1.30, 2.30, 2.85] if len(rows[0]) == 3 else [6.45/len(rows[0])]*len(rows[0])
            for row_index, row in enumerate(rows):
                for cell_index, value in enumerate(row):
                    cell = table.cell(row_index, cell_index)
                    cell.width = Inches(widths[cell_index])
                    inline(cell.paragraphs[0], value)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(10)
                        if row_index == 0:
                            run.bold = True
                properties = table.rows[row_index]._tr.get_or_add_trPr()
                no_split = OxmlElement("w:cantSplit"); properties.append(no_split)
                if row_index == 0:
                    repeat = OxmlElement("w:tblHeader"); properties.append(repeat)
            document.add_paragraph()
            table_count += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title"); inline(paragraph, line[2:])
        elif line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1"); inline(paragraph, line[3:])
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2"); inline(paragraph, line[4:])
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet"); inline(paragraph, line[2:])
        else:
            parts = [line]
            while index + 1 < len(lines) and lines[index + 1].strip() and not lines[index + 1].lstrip().startswith(("#", "|", "- ")):
                index += 1
                parts.append(lines[index].strip())
            paragraph = document.add_paragraph(); inline(paragraph, " ".join(parts))
        paragraph_count += 1
        index += 1
    document.core_properties.title = "Adversarial discovery of neural mechanisms in human consciousness"
    document.core_properties.subject = "Clinical-model-guided human discovery, competing mechanisms and independent diagnostic tests"
    document.save(path)
    print({"paragraphs": paragraph_count, "tables": table_count, "source": str(root / "Adversarial_Consciousness_Research_Proposal.md"), "output": str(path), "preserved_original": str(backup)})


if __name__ == "__main__":
    update(Path(__file__).resolve().parents[2])
